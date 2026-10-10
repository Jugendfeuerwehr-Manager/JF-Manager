"""Exports and qualification attachments follow scoped model rights."""

from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from departments.models import Department, UserDepartmentRole
from members.models import Member, Parent
from qualifications.models import Qualification, QualificationType, SpecialTask, SpecialTaskType


class ExportAndQualificationAttachmentRoleTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.department = Department.objects.create(name="A", code="export-qualification-a")
        cls.member = Member.objects.create(name="A", lastname="Member")
        cls.member.departments.add(cls.department)
        qualification_type = QualificationType.objects.create(name="Course")
        cls.qualification = Qualification.objects.create(
            member=cls.member, type=qualification_type, date_acquired=date(2026, 1, 1)
        )
        task_type = SpecialTaskType.objects.create(name="Duty")
        cls.special_task = SpecialTask.objects.create(member=cls.member, task=task_type, start_date=date(2026, 1, 1))
        cls.user = get_user_model().objects.create_user(username="qualified-viewer", password="test-only-password")
        group = Group.objects.create(name="Scoped member and qualification editor")
        group.permissions.add(
            Permission.objects.get(content_type__app_label="members", codename="view_member"),
            Permission.objects.get(content_type__app_label="members", codename="export_member"),
            Permission.objects.get(content_type__app_label="qualifications", codename="view_qualification"),
            Permission.objects.get(content_type__app_label="qualifications", codename="add_qualification"),
            Permission.objects.get(content_type__app_label="qualifications", codename="change_qualification"),
            Permission.objects.get(content_type__app_label="qualifications", codename="view_specialtask"),
            Permission.objects.get(content_type__app_label="qualifications", codename="change_specialtask"),
        )
        UserDepartmentRole.objects.create(user=cls.user, department=cls.department).groups.add(group)

    def setUp(self):
        self.client.force_authenticate(user=self.user)

    def test_scoped_member_export_right_allows_scoped_export(self):
        department_b = Department.objects.create(name="B", code="export-qualification-b")
        member_b = Member.objects.create(name="B", lastname="Member")
        member_b.departments.add(department_b)
        response = self.client.get("/api/v1/members/export-excel/?columns=name")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        sheet = load_workbook(BytesIO(response.content), read_only=True).active
        self.assertEqual([row[0] for row in list(sheet.values)[1:]], ["A"])

    def test_export_does_not_include_parents_without_parent_view_right(self):
        self.user.user_permissions.add(
            Permission.objects.get(content_type__app_label="members", codename="view_member")
        )
        parent = Parent.objects.create(name="Parent", lastname="Example", email="private-parent@example.invalid")
        parent.children.add(self.member)

        response = self.client.get("/api/v1/members/export-excel/?columns=name,parent1_email")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        sheet = load_workbook(BytesIO(response.content), read_only=True).active
        self.assertEqual(sheet.cell(row=2, column=2).value, None)

        role = UserDepartmentRole.objects.get(user=self.user, department=self.department)
        role.groups.first().permissions.add(
            Permission.objects.get(content_type__app_label="members", codename="view_parent")
        )
        allowed_response = self.client.get("/api/v1/members/export-excel/?columns=name,parent1_email")
        allowed_sheet = load_workbook(BytesIO(allowed_response.content), read_only=True).active
        self.assertEqual(allowed_sheet.cell(row=2, column=2).value, "private-parent@example.invalid")

    def test_scoped_qualification_edit_right_allows_attachment_upload(self):
        upload = SimpleUploadedFile("course.pdf", b"%PDF-1.4\n", content_type="application/pdf")
        response = self.client.post(
            f"/api/v1/qualifications/{self.qualification.pk}/attachments/",
            {"file": upload, "name": "Course document"},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(self.qualification.attachments.count(), 1)
        preview = self.client.get(response.data["file_url"])
        self.assertEqual(preview.status_code, 200)
        b"".join(preview.streaming_content)  # closes the file without request_finished

        deleted = self.client.delete(
            f"/api/v1/qualifications/{self.qualification.pk}/attachments/{response.data['id']}/"
        )
        self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(self.qualification.attachments.count(), 0)

    def test_scoped_special_task_edit_right_allows_attachment_upload(self):
        upload = SimpleUploadedFile("duty.pdf", b"%PDF-1.4\n", content_type="application/pdf")
        response = self.client.post(
            f"/api/v1/qualifications/specialtasks/{self.special_task.pk}/attachments/",
            {"file": upload, "name": "Duty document"},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(self.special_task.attachments.count(), 1)
