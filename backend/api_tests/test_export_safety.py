from datetime import date, timedelta
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.management import call_command
from django.test import SimpleTestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from departments.models import Department, UserDepartmentRole
from jf_manager_backend.safe_exports import append_safe_row
from members.models import ExportAudit, Member, MemberList, MemberListEntry, Parent
from members.renderers import MemberExcelRenderer


class ExportCellsTests(SimpleTestCase):
    def test_strings_are_not_formulas_and_dates_and_numbers_keep_types(self):
        workbook = Workbook()
        values = ['=HYPERLINK("https://example.test")', "+SUM(1,2)", "-1", "@SUM(1)", "00123", 12, date(2030, 1, 2)]
        append_safe_row(workbook.active, values)
        output = BytesIO()
        workbook.save(output)
        cells = next(load_workbook(BytesIO(output.getvalue())).active.rows)
        self.assertEqual([cell.data_type for cell in cells], ["s", "s", "s", "s", "s", "n", "d"])
        self.assertEqual(cells[0].value, values[0])

    def test_legacy_renderer_uses_the_same_text_contract(self):
        data = dict.fromkeys(
            [
                "name",
                "lastname",
                "email",
                "street",
                "zip_code",
                "city",
                "phone",
                "mobile",
                "notes",
                "identityCardNumber",
                "status",
            ],
            "=1+1",
        )
        data.update(birthday="2030-01-02", joined=None, canSwimm=False, parents=[])
        workbook = load_workbook(BytesIO(MemberExcelRenderer().render([data])))
        self.assertEqual(workbook.active.cell(2, 1).data_type, "s")
        self.assertEqual(workbook.active.cell(2, 3).data_type, "d")


class ExportBoundaryTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="export-reader")
        self.a = Department.objects.create(name="A", code="export-a")
        self.b = Department.objects.create(name="B", code="export-b")
        self.read = Group.objects.create(name="Export read only")
        self.read.permissions.add(Permission.objects.get(codename="view_member"))
        self.export = Group.objects.create(name="Export allowed")
        self.export.permissions.add(
            Permission.objects.get(codename="export_member"),
            Permission.objects.get(codename="view_memberlist"),
            Permission.objects.get(codename="export_memberlist"),
        )
        role = UserDepartmentRole.objects.create(user=self.user, department=self.a)
        role.groups.add(self.read)
        self.role = role
        UserDepartmentRole.objects.create(user=self.user, department=self.b).groups.add(self.read)
        self.member = Member.objects.create(
            name="=1+1", lastname="Synthetic", birthday=date(2012, 2, 3), zip_code="00123"
        )
        self.member.departments.add(self.a)
        other = Member.objects.create(name="Foreign", lastname="Synthetic")
        other.departments.add(self.b)
        self.client.force_authenticate(self.user)

    def test_export_requires_extra_right_and_audits_denial_and_success(self):
        url = "/api/v1/members/export-excel/?columns=name,birthday,age,zip_code"
        self.assertEqual(self.client.get(url).status_code, 403)
        self.role.groups.add(self.export)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([sheet.cell(2, column).data_type for column in range(1, 5)], ["s", "d", "n", "s"])
        self.assertEqual(sheet.cell(2, 4).value, "00123")
        self.assertEqual(list(ExportAudit.objects.order_by("id").values_list("status_code", flat=True)), [403, 200])
        self.assertEqual(ExportAudit.objects.first().actor_id, self.user.pk)
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_list_export_uses_text_cells_and_does_not_bypass_parent_permission(self):
        self.role.groups.add(self.export)
        listing = MemberList.objects.create(name="Synthetic list", department=self.a)
        MemberListEntry.objects.create(member_list=listing, member=self.member, notes="=1+1")
        parent = Parent.objects.create(name="Hidden", lastname="Parent", email="synthetic@example.test")
        parent.children.add(self.member)
        response = self.client.get(
            f"/api/v1/member-lists/{listing.pk}/export-excel/?columns=name,list_notes,parent1_email"
        )
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertIsNone(sheet.cell(2, 3).value)

    def test_list_export_accepts_names_with_excel_reserved_characters(self):
        self.role.groups.add(self.export)
        listing = MemberList.objects.create(name="Ausflug / Anmeldung: [Gruppe A]?*\\", department=self.a)
        MemberListEntry.objects.create(member_list=listing, member=self.member, checked=True, checked_at=timezone.now())
        response = self.client.get(f"/api/v1/member-lists/{listing.pk}/export-excel/?columns=name,list_checked_at")
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertTrue(sheet.title)
        self.assertLessEqual(len(sheet.title), 31)
        self.assertFalse(any(char in sheet.title for char in "\\/*?:[]"))
        self.assertEqual(sheet.cell(2, 1).data_type, "s")
        self.assertEqual(sheet.cell(2, 2).data_type, "d")
        self.assertEqual(response["Cache-Control"], "private, no-store")

    @override_settings(AUDIT_RETENTION_DAYS=180)
    def test_audit_retention_removes_only_expired_metadata(self):
        old = ExportAudit.objects.create(actor_id=self.user.pk, object_type="members.member", status_code=200)
        ExportAudit.objects.filter(pk=old.pk).update(created_at=timezone.now() - timedelta(days=181))
        recent = ExportAudit.objects.create(actor_id=self.user.pk, object_type="members.member", status_code=403)
        call_command("purge_export_audits", stdout=StringIO())
        self.assertEqual(list(ExportAudit.objects.values_list("pk", flat=True)), [recent.pk])
